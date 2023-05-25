from ..document import Document
from .block import *
from openpyxl import load_workbook


class ExcelDocument(Document):
    def __init__(self, path):
        super().__init__(path)

    def _parse_fields(self):
        # TODO: manage other table types

        fields = {}

        for sheet in self._f_instance.worksheets:
            # print("sheet " + sheet.title)
            headers = {}
            letter = "A"
            while sheet[f"{letter}1"].value is not None:
                headers[f"{letter}"] = sheet[f"{letter}1"].value
                letter = next_letter(letter)

            # print("\t", headers)

            rows = {}
            row_ind = 2
            while sheet[f"A{row_ind}"].value is not None:
                cells = {}
                for letter, title in headers.items():
                    cells[title] = ExcelCell(sheet[f"{letter}{row_ind}"], f"{letter}{row_ind}")
                rows[str(row_ind - 2)] = ExcelRow(None, row_ind - 2, cells)

                row_ind += 1
                # print("\tadding row", row_ind - 2)

            fields[sheet.title] = ExcelSheet(sheet, sheet.title, rows, headers)

        return fields

    def open(self):
        # print(self._path)
        return load_workbook(self._path)

    def save(self, path):
        self._f_instance.save(path)
